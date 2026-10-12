from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from zoneinfo import ZoneInfo

from .config import app_home
from .dashboard import render_dashboard
from .database import Database


PALETTE = {"bg": "101519", "panel": "181F24", "panel2": "1E272D", "line": "303C44", "text": "EEF2F3", "muted": "A4B3BC", "mint": "7ED8BC", "blue": "88BBED", "amber": "E4BF73", "red": "EE9292"}
STATUS_LABELS = {"active": "Activa", "unknown": "Por verificar", "discarded": "Descartada", "rejected": "Rechazada", "offer": "Oferta", "hired": "Contratada"}
STAGE_LABELS = {"unconfirmed": "Sin confirmar", "application_received": "Recibida", "screening": "Filtro inicial", "assessment": "Prueba / evaluación", "interview": "Entrevista", "reference": "Referencias", "offer": "Oferta", "hired": "Contratada", "rejected": "Rechazada"}


def _safe(value: Any) -> str:
    return "" if value is None else str(value)


def _style_sheet(ws, header_row: int, widths: dict[int, int]) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=PALETTE["panel2"])
            cell.font = Font(name="Aptos", size=10, bold=True, color=PALETTE["mint"])
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=PALETTE["line"]))
    ws.row_dimensions[header_row].height = 28
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color=PALETTE["text"])
            cell.fill = PatternFill("solid", fgColor=PALETTE["bg"] if cell.row % 2 else PALETTE["panel"])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=PALETTE["line"]))


def build_workbook(data: dict[str, Any], output: Path) -> Path:
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_setup.orientation = "landscape"
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    dashboard.sheet_view.zoomScale = 85
    for row in range(1, 36):
        for col in range(1, 13):
            c = dashboard.cell(row, col)
            c.fill = PatternFill("solid", fgColor=PALETTE["bg"])
    for col in range(1, 13):
        dashboard.column_dimensions[get_column_letter(col)].width = 14
    dashboard.merge_cells("A1:L2")
    dashboard["A1"] = "APLICACIÓN A EMPLEOS  /  SEGUIMIENTO LABORAL"
    dashboard["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=PALETTE["text"])
    dashboard["A1"].fill = PatternFill("solid", fgColor=PALETTE["panel"])
    dashboard["A1"].alignment = Alignment(vertical="center", indent=1)
    dashboard.merge_cells("A3:L3")
    dashboard["A3"] = f"Corte generado {data.get('generated_at','')}  ·  Afinidad y avance se presentan por separado"
    dashboard["A3"].font = Font(name="Aptos", size=10, color=PALETTE["muted"])

    apps = data.get("applications", [])
    company_map = {a["id"]: a["company"] for a in apps}
    active = sum(a["status"] in {"active", "unknown", "offer", "hired"} for a in apps)
    interviews = sum(a["stage"] == "interview" for a in apps)
    promising = sum(a.get("affinity_score") is not None and float(a["affinity_score"]) >= 80 for a in apps)
    review_app_ids = {r.get("application_id") for r in data.get("reviews", []) if r.get("application_id") is not None}
    pending_reviews = len(data.get("reviews", [])) + sum(bool(a.get("needs_review")) and a.get("id") not in review_app_ids for a in apps)
    metrics = [("A5:C5", "A6:C7", "CANDIDATURAS ACTIVAS", active, PALETTE["mint"]),
               ("D5:F5", "D6:F7", "ENTREVISTAS", interviews, PALETTE["blue"]),
               ("G5:I5", "G6:I7", "PROMETEDORAS", promising, PALETTE["amber"]),
               ("J5:L5", "J6:L7", "POR REVISAR", pending_reviews, PALETTE["red"])]
    for label_range, value_range, label, value, accent in metrics:
        dashboard.merge_cells(label_range)
        dashboard.merge_cells(value_range)
        label_cell = dashboard[label_range.split(":")[0]]
        label_cell.value = label
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=PALETTE["muted"])
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell = dashboard[value_range.split(":")[0]]
        value_cell.value = value
        value_cell.font = Font(name="Aptos Display", size=24, bold=True, color=accent)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in range(int(label_range[1]), int(value_range.split(":")[1][1]) + 1):
            for col in range(ord(label_range[0]) - 64, ord(value_range.split(":")[1][0]) - 64 + 1):
                dashboard.cell(row, col).fill = PatternFill("solid", fgColor=PALETTE["panel"])
    dashboard.row_dimensions[5].height = 21
    dashboard.row_dimensions[6].height = 26
    dashboard.row_dimensions[7].height = 17

    dashboard["A10"] = "ETAPA"
    dashboard["B10"] = "CANDIDATURAS"
    stage_counts: dict[str, int] = {}
    for app in apps:
        label = STAGE_LABELS.get(app.get("stage", ""), app.get("stage", ""))
        stage_counts[label] = stage_counts.get(label, 0) + 1
    for idx, (stage, count) in enumerate(sorted(stage_counts.items(), key=lambda item: (-item[1], item[0])) or [("Sin registros", 0)], start=11):
        dashboard.cell(idx, 1, stage)
        dashboard.cell(idx, 2, count)
    dashboard["D10"] = "ESTADO"
    dashboard["E10"] = "CANDIDATURAS"
    state_counts: dict[str, int] = {}
    for app in apps:
        key = "discarded" if app.get("status") == "rejected" else app.get("status", "unknown")
        label = STATUS_LABELS.get(key, key)
        state_counts[label] = state_counts.get(label, 0) + 1
    for idx, (status, count) in enumerate(sorted(state_counts.items(), key=lambda item: (-item[1], item[0])) or [("Sin registros", 0)], start=11):
        dashboard.cell(idx, 4, status)
        dashboard.cell(idx, 5, count)
    for coord in ("A10", "B10", "D10", "E10"):
        dashboard[coord].font = Font(name="Aptos", bold=True, size=9, color=PALETTE["mint"])

    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Candidaturas por etapa"
    bar.y_axis.title = "Etapa"
    bar.x_axis.title = "Cantidad"
    bar.height = 7.3
    bar.width = 13.8
    bar.legend = None
    bar.add_data(Reference(dashboard, min_col=2, min_row=10, max_row=10 + max(1, len(stage_counts))), titles_from_data=True)
    bar.set_categories(Reference(dashboard, min_col=1, min_row=11, max_row=10 + max(1, len(stage_counts))))
    bar.series[0].graphicalProperties.solidFill = PALETTE["blue"]
    bar.series[0].graphicalProperties.line.solidFill = PALETTE["blue"]
    dashboard.add_chart(bar, "A19")

    donut = DoughnutChart()
    donut.title = "Estado de los procesos"
    donut.holeSize = 62
    donut.height = 7.3
    donut.width = 12.8
    donut.add_data(Reference(dashboard, min_col=5, min_row=10, max_row=10 + max(1, len(state_counts))), titles_from_data=True)
    donut.set_categories(Reference(dashboard, min_col=4, min_row=11, max_row=10 + max(1, len(state_counts))))
    donut.dataLabels = DataLabelList()
    donut.dataLabels.showPercent = True
    dashboard.add_chart(donut, "G19")

    applications = wb.create_sheet("Candidaturas")
    headers = ["Empresa", "Cargo", "Estado", "Etapa", "Afinidad %", "Etiqueta", "Última novedad", "Próxima acción", "Fecha objetivo", "Revisión pendiente", "Origen"]
    applications.append(headers)
    for item in apps:
        applications.append([
            item.get("company", ""), item.get("role", ""), STATUS_LABELS.get(item.get("status", ""), item.get("status", "")),
            STAGE_LABELS.get(item.get("stage", ""), item.get("stage", "")), item.get("affinity_score"), item.get("affinity_label", ""),
            item.get("last_update", ""), item.get("next_action", ""), item.get("next_action_due", ""),
            "Sí" if item.get("needs_review") else "No", item.get("source", ""),
        ])
    _style_sheet(applications, 1, {1: 24, 2: 29, 3: 17, 4: 19, 5: 13, 6: 20, 7: 25, 8: 48, 9: 18, 10: 17, 11: 20})
    applications.auto_filter.ref = applications.dimensions

    history = wb.create_sheet("Historial")
    history.append(["Fecha", "Empresa", "Tipo de novedad", "Etapa", "Estado", "Remitente", "Asunto", "Resumen de evidencia", "Fuente", "Confianza", "Revisión"])
    company_map = {a["id"]: a["company"] for a in apps}
    for event in data.get("events", []):
        history.append([
            event.get("occurred_at", ""), company_map.get(event.get("application_id"), ""), event.get("event_type", ""),
            STAGE_LABELS.get(event.get("stage", ""), event.get("stage", "")), STATUS_LABELS.get(event.get("status", ""), event.get("status", "")),
            event.get("sender", ""), event.get("subject", ""), event.get("evidence_summary", ""), event.get("source_url", ""), event.get("confidence", 0),
            "Sí" if event.get("requires_review") else "No",
        ])
    _style_sheet(history, 1, {1: 25, 2: 22, 3: 23, 4: 20, 5: 16, 6: 32, 7: 40, 8: 70, 9: 50, 10: 12, 11: 12})
    history.auto_filter.ref = history.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def build_pdf(data: dict[str, Any], output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output), pagesize=landscape(letter), rightMargin=0.42*inch, leftMargin=0.42*inch, topMargin=0.45*inch, bottomMargin=0.45*inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=19, leading=22, textColor=colors.HexColor("#172b36"), alignment=TA_LEFT, spaceAfter=4))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#465f6e")))
    styles.add(ParagraphStyle(name="Cell", parent=styles["BodyText"], fontSize=7.2, leading=9, textColor=colors.HexColor("#172b36")))
    story = [Paragraph("Aplicación a Empleos · Informe de candidaturas", styles["ReportTitle"]),
             Paragraph(f"Corte: {data.get('generated_at','')} · Afinidad estimada y avance del proceso se informan por separado.", styles["Small"]), Spacer(1, 0.14*inch)]
    apps = data.get("applications", [])
    company_map = {a["id"]: a["company"] for a in apps}
    active = sum(a["status"] in {"active", "unknown", "offer", "hired"} for a in apps)
    interviews = sum(a["stage"] == "interview" for a in apps)
    promising = sum(a.get("affinity_score") is not None and float(a["affinity_score"]) >= 80 for a in apps)
    review_app_ids = {r.get("application_id") for r in data.get("reviews", []) if r.get("application_id") is not None}
    review_count = len(data.get("reviews", [])) + sum(bool(a.get("needs_review")) and a.get("id") not in review_app_ids for a in apps)
    metrics = [["Candidaturas activas", "Entrevistas", "Prometedoras", "Por revisar"], [str(active), str(interviews), str(promising), str(review_count)]]
    metric_table = Table(metrics, colWidths=[2.45*inch]*4, rowHeights=[0.28*inch,0.4*inch])
    metric_table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#eef3f6")),("TEXTCOLOR",(0,0),(-1,0),colors.HexColor("#4a6575")),
        ("TEXTCOLOR",(0,1),(0,1),colors.HexColor("#08765a")),("FONTNAME",(0,0),(-1,0),"Helvetica"),
        ("FONTNAME",(0,1),(-1,1),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),8),("FONTSIZE",(0,1),(-1,1),16),
        ("ALIGN",(0,0),(-1,-1),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOX",(0,0),(-1,-1),0.5,colors.HexColor("#ccd8df")),
        ("INNERGRID",(0,0),(-1,-1),0.4,colors.HexColor("#d6e0e5")),
    ]))
    story.extend([metric_table, Spacer(1,0.15*inch), Paragraph("Candidaturas", styles["Heading2"])])
    table_data = [["Empresa / cargo", "Estado", "Etapa", "Afinidad", "Última novedad", "Próxima acción"]]
    for item in apps:
        company_role = _safe(item.get("company")) + (f" / {_safe(item.get('role'))}" if item.get("role") else "")
        score = "—" if item.get("affinity_score") is None else f"{item['affinity_score']}% · {item.get('affinity_label','')}"
        table_data.append([
            Paragraph(company_role[:95], styles["Cell"]),
            Paragraph(STATUS_LABELS.get(item.get("status", ""), _safe(item.get("status"))), styles["Cell"]),
            Paragraph(STAGE_LABELS.get(item.get("stage", ""), _safe(item.get("stage"))), styles["Cell"]),
            Paragraph(score, styles["Cell"]),
            Paragraph(_safe(item.get("last_update"))[:32] or "—", styles["Cell"]),
            Paragraph(_safe(item.get("next_action"))[:120], styles["Cell"]),
        ])
    if len(table_data) == 1:
        table_data.append([Paragraph("No hay candidaturas cargadas.", styles["Cell"]), "", "", "", "", ""])
    app_table = Table(table_data, repeatRows=1, colWidths=[2.0*inch,1.0*inch,1.25*inch,1.25*inch,1.8*inch,3.15*inch])
    app_table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#18242c")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),7),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#cfdae0")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f3f6f8")]),
        ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),
    ]))
    story.extend([app_table, Spacer(1,0.13*inch), Paragraph("Novedades recientes", styles["Heading2"])])
    event_rows = [["Fecha", "Empresa", "Novedad", "Asunto / resumen"]]
    for event in data.get("events", [])[:12]:
        event_rows.append([_safe(event.get("occurred_at"))[:25], company_map.get(event.get("application_id"), ""), _safe(event.get("event_type")), Paragraph((_safe(event.get("subject")) + " — " + _safe(event.get("evidence_summary")))[:230], styles["Cell"])])
    if len(event_rows) == 1:
        event_rows.append(["—", "—", "Sin novedades confirmadas", "Los correos completos permanecen en la cuenta de origen."])
    event_table = Table(event_rows, repeatRows=1, colWidths=[1.75*inch,1.8*inch,1.8*inch,5.1*inch])
    event_table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1e272d")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),7),
        ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#cfdae0")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f3f6f8")]),
        ("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,1),(-1,-1),7),
    ]))
    story.append(event_table)
    if data.get("reviews"):
        story.extend([Spacer(1,0.13*inch), Paragraph("Revisión pendiente", styles["Heading2"])])
        review_rows = [["Empresa sugerida", "Fecha", "Asunto", "Motivo"]]
        for item in data["reviews"][:10]:
            review_rows.append([_safe(item.get("company_guess")) or "Por identificar", _safe(item.get("occurred_at"))[:25], Paragraph(_safe(item.get("subject"))[:100], styles["Cell"]), Paragraph(_safe(item.get("reason")), styles["Cell"])])
        review_table = Table(review_rows, repeatRows=1, colWidths=[1.7*inch,1.7*inch,3.3*inch,3.75*inch])
        review_table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#302b1f")),("TEXTCOLOR",(0,0),(-1,0),colors.HexColor("#704f0d")),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#cfdae0")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f3f6f8")]),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),7)]))
        story.append(review_table)
    doc.build(story)
    return output


def generate_reports(database: Database, output_dir: Path | None = None,
                    timezone_name: str = "America/Bogota", run_label: str = "") -> dict[str, Path]:
    data = database.dashboard_data()
    output_dir = output_dir or (app_home() / "reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now(ZoneInfo(timezone_name))
    data["generated_at"] = generated_at.isoformat(timespec="seconds")
    stamp = generated_at.strftime("%Y%m%d_%H%M%S")
    if run_label:
        stamp += "_" + "".join(character for character in run_label if character.isalnum() or character in "-_")
    xlsx = build_workbook(data, output_dir / f"candidaturas_{stamp}.xlsx")
    pdf = build_pdf(data, output_dir / f"candidaturas_{stamp}.pdf")
    html = output_dir / "dashboard.html"
    html.write_bytes(render_dashboard(data))
    dashboard_xlsx = output_dir / "dashboard.xlsx"
    shutil.copyfile(xlsx, dashboard_xlsx)
    return {"xlsx": xlsx, "pdf": pdf, "html": html, "dashboard_xlsx": dashboard_xlsx}
