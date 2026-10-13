from __future__ import annotations

import base64
import email
import json
import unittest
import uuid
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote
from unittest.mock import Mock, patch
from email import policy

from openpyxl import load_workbook

from job_tracker.cloud import GoogleDrivePublisher
from job_tracker.config import CloudSettings, MailboxSettings
from job_tracker.dashboard import render_dashboard
from job_tracker.database import Database
from job_tracker.emailer import EmailDeliveryError, GmailReportSender
from job_tracker.mailbox import GmailApiMailbox, ImapMailbox, parse_rfc822
from job_tracker.models import Application, EmailMessage
from job_tracker.oauth import GoogleOAuthClient
from job_tracker.reporting import build_pdf, build_workbook, generate_reports
from job_tracker.security import CredentialStore, CredentialStoreError


TEST_TMP = Path(__file__).resolve().parents[1] / ".local" / "tmp"


def temporary_path(name: str, suffix: str = "") -> Path:
    TEST_TMP.mkdir(parents=True, exist_ok=True)
    return TEST_TMP / f"{name}-{uuid.uuid4().hex}{suffix}"


class MemoryStore:
    def __init__(self):
        self.values = {}

    def put(self, key, secret):
        self.values[key] = secret

    def get(self, key):
        return self.values.get(key)


class MailboxTests(unittest.TestCase):
    def test_parse_rfc822_headers_and_text(self):
        raw = (b"From: Recruiting <jobs@example.com>\r\nSubject: =?utf-8?q?Entrevista?=\r\n"
               b"Date: Mon, 12 Oct 2026 10:00:00 -0500\r\nMessage-ID: <abc@example.com>\r\n"
               b"Content-Type: text/plain; charset=utf-8\r\n\r\nTe invitamos a entrevista.")
        message = parse_rfc822("imap", "uid-1", raw)
        self.assertEqual(message.subject, "Entrevista")
        self.assertEqual(message.message_id, "abc@example.com")
        self.assertIn("entrevista", message.body_text.casefold())

    def test_imap_select_is_readonly_and_fetch_uses_peek(self):
        raw = (b"From: Recruiting <jobs@example.com>\r\nSubject: Entrevista\r\n"
               b"Date: Mon, 12 Oct 2026 10:00:00 -0500\r\nMessage-ID: <abc>\r\n"
               b"Content-Type: text/plain; charset=utf-8\r\n\r\nEntrevista.")
        fake = Mock()
        fake.login.return_value = ("OK", [])
        fake.select.return_value = ("OK", [b"1"])
        fake.uid.side_effect = [("OK", [b"1"]), ("OK", [(b"1 BODY[]", raw), b")"])]
        fake.logout.return_value = ("BYE", [])
        settings = MailboxSettings(provider="imap", email="test@example.com", imap_host="imap.example.com")
        with patch("job_tracker.mailbox.imaplib.IMAP4_SSL", return_value=fake):
            results = list(ImapMailbox(settings, "dummy-app-password").fetch(date(2026, 10, 1)))
        fake.select.assert_called_once_with("INBOX", readonly=True)
        self.assertIn("BODY.PEEK[]", fake.uid.call_args_list[1].args[-1])
        self.assertEqual(len(results), 1)

    def test_gmail_api_query_is_readonly_and_parses_message(self):
        oauth = Mock()
        oauth.token.return_value = "synthetic-access-token"
        mailbox = GmailApiMailbox(oauth, "newer_than:365d", "test@example.com")
        body = base64.urlsafe_b64encode(b"Hemos recibido tu candidatura.").decode().rstrip("=")
        result = {"messages": [{"id": "synthetic-id"}]}
        details = {"id": "synthetic-id", "internalDate": "1791817200000", "payload": {
            "mimeType": "text/plain", "headers": [
                {"name": "Subject", "value": "Application received"},
                {"name": "From", "value": "Hiring <jobs@example.com>"},
                {"name": "Message-ID", "value": "<synthetic@example.com>"},
            ], "body": {"data": body},
        }}
        with patch.object(mailbox, "_get_json", side_effect=[result, details]) as get_json:
            messages = list(mailbox.fetch(date(2026, 10, 1)))
        self.assertEqual(len(messages), 1)
        self.assertIn("after:2026/10/01", unquote(get_json.call_args_list[0].args[0]))
        self.assertEqual(messages[0].source, "gmail:test@example.com")
        self.assertIn("candidatura", messages[0].body_text)
        self.assertTrue(messages[0].source_url.startswith("https://mail.google.com/"))


class SecurityAndCloudTests(unittest.TestCase):
    def test_google_mail_and_drive_use_separate_narrow_scopes(self):
        self.assertEqual(GoogleOAuthClient.GMAIL_SCOPES, ("https://www.googleapis.com/auth/gmail.readonly",))
        self.assertEqual(GoogleOAuthClient.GMAIL_SEND_SCOPES, (
            "https://www.googleapis.com/auth/gmail.readonly",
            "https://www.googleapis.com/auth/gmail.send",
        ))
        self.assertEqual(GoogleOAuthClient.DRIVE_SCOPES, ("https://www.googleapis.com/auth/drive.file",))

    def test_oauth_client_file_is_local_input_not_embedded_in_settings(self):
        path = temporary_path("client", ".json")
        try:
            path.write_text('{"installed":{"client_id":"local-client","client_secret":"local-secret"}}', encoding="utf-8")
            store = MemoryStore()
            client = GoogleOAuthClient(path, store, "user@example.com", GoogleOAuthClient.GMAIL_SCOPES, "gmail")
        finally:
            path.unlink(missing_ok=True)
        self.assertEqual(client.client["client_id"], "local-client")
        self.assertEqual(store.values, {})

    def test_non_windows_secret_storage_fails_closed(self):
        store = CredentialStore()
        if __import__("os").name != "nt":
            with self.assertRaises(CredentialStoreError):
                store.put("mailbox-password", "secret")

    def test_google_multipart_upload_contains_metadata_and_media(self):
        body, content_type = GoogleDrivePublisher._multipart({"name": "report.xlsx"}, b"synthetic-file", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("multipart/related", content_type)
        self.assertIn(b"report.xlsx", body)
        self.assertIn(b"synthetic-file", body)


class GmailReportTests(unittest.TestCase):
    def setUp(self):
        self.client_path = temporary_path("gmail-client", ".json")
        self.client_path.write_text('{"installed":{"client_id":"synthetic-client"}}', encoding="utf-8")
        self.pdf_path = temporary_path("daily-report", ".pdf")
        self.xlsx_path = temporary_path("daily-report", ".xlsx")
        self.pdf_path.write_bytes(b"%PDF-synthetic")
        self.xlsx_path.write_bytes(b"synthetic-xlsx")
        self.store = MemoryStore()

    def tearDown(self):
        for path in (self.client_path, self.pdf_path, self.xlsx_path):
            path.unlink(missing_ok=True)

    def test_sends_only_to_authenticated_account_with_pdf_and_xlsx(self):
        sender = GmailReportSender(
            MailboxSettings(provider="gmail", gmail_client_file=str(self.client_path)), self.store
        )
        sender.oauth.token = Mock(return_value="synthetic-access-token")
        requests = []

        class Response:
            def __init__(self, value):
                self.value = value
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False
            def read(self):
                return json.dumps(self.value).encode("utf-8")

        def fake_urlopen(request, timeout):
            requests.append(request)
            if request.full_url.endswith("/profile"):
                return Response({"emailAddress": "self@example.test"})
            return Response({"id": "synthetic-sent-id"})

        with patch("job_tracker.emailer.urllib.request.urlopen", side_effect=fake_urlopen):
            result = sender.send(
                {"pdf": self.pdf_path, "xlsx": self.xlsx_path},
                "morning", datetime(2026, 10, 13, 8, 0),
            )

        self.assertEqual(result, {"status": "sent", "attachment_count": 2})
        self.assertEqual(len(requests), 2)
        self.assertTrue(requests[0].full_url.endswith("/profile"))
        self.assertEqual(requests[1].get_method(), "POST")
        sent = json.loads(requests[1].data)
        raw = base64.urlsafe_b64decode(sent["raw"] + "=" * (-len(sent["raw"]) % 4))
        parsed = email.message_from_bytes(raw, policy=policy.default)
        self.assertEqual(parsed["To"], "self@example.test")
        self.assertEqual(parsed["From"], "self@example.test")
        self.assertEqual({part.get_filename() for part in parsed.iter_attachments()},
                         {self.pdf_path.name, self.xlsx_path.name})

    def test_refuses_to_send_if_authenticated_account_differs_from_configured_mailbox(self):
        sender = GmailReportSender(
            MailboxSettings(provider="gmail", email="expected@example.test",
                            gmail_client_file=str(self.client_path)), self.store
        )
        sender.oauth.token = Mock(return_value="synthetic-access-token")
        response = Mock()
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=False)
        response.read.return_value = b'{"emailAddress":"other@example.test"}'
        with patch("job_tracker.emailer.urllib.request.urlopen", return_value=response) as request:
            with self.assertRaises(EmailDeliveryError):
                sender.send({"pdf": self.pdf_path, "xlsx": self.xlsx_path}, "morning", datetime.now())
        request.assert_called_once()

    def test_refuses_to_send_if_authenticated_account_differs_from_cloud_account(self):
        sender = GmailReportSender(
            MailboxSettings(provider="gmail", gmail_client_file=str(self.client_path)), self.store,
            expected_account="expected@example.test",
        )
        sender.oauth.token = Mock(return_value="synthetic-access-token")
        response = Mock()
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=False)
        response.read.return_value = b'{"emailAddress":"other@example.test"}'
        with patch("job_tracker.emailer.urllib.request.urlopen", return_value=response) as request:
            with self.assertRaises(EmailDeliveryError):
                sender.send({"pdf": self.pdf_path, "xlsx": self.xlsx_path}, "morning", datetime.now())
        request.assert_called_once()


class ArtifactTests(unittest.TestCase):
    def test_reports_generate_dashboard_excel_pdf_and_standalone_html(self):
        db_path = temporary_path("tracker", ".sqlite3")
        output_dir = TEST_TMP
        db = Database(db_path)
        try:
            app_id = db.add_application(Application(None, "Synthetic Corp", "Analyst", status="active", stage="interview", affinity_score=85, affinity_label="promising", needs_review=False, last_update="2026-10-12T09:00:00-05:00"))
            app = db.get_application(app_id)
            data = db.dashboard_data()
            html = render_dashboard({**data, "applications": [{**app, "company": "</script><script>alert(1)</script>"}]})
            self.assertNotIn(b"</script><script>alert(1)", html)
            files = generate_reports(db, output_dir)
            self.assertTrue(files["pdf"].read_bytes().startswith(b"%PDF"))
            self.assertNotIn(b"__TRACKER_DATA__", files["html"].read_bytes())
            workbook = load_workbook(files["dashboard_xlsx"])
            self.assertIn("Dashboard", workbook.sheetnames)
            self.assertIn("Candidaturas", workbook.sheetnames)
            self.assertIn("Historial", workbook.sheetnames)
            self.assertGreaterEqual(len(workbook["Dashboard"]._charts), 2)
            self.assertEqual(workbook["Candidaturas"].cell(2, 1).value, "Synthetic Corp")
        finally:
            db_path.unlink(missing_ok=True)
            Path(f"{db_path}-wal").unlink(missing_ok=True)
            Path(f"{db_path}-shm").unlink(missing_ok=True)
            for artifact in locals().get("files", {}).values():
                artifact.unlink(missing_ok=True)

    def test_pdf_and_xlsx_empty_dataset(self):
        pdf_path = temporary_path("empty", ".pdf")
        xlsx_path = temporary_path("empty", ".xlsx")
        empty = {"generated_at": "2026-10-12T18:00:00-05:00", "applications": [], "events": [], "reviews": []}
        build_pdf(empty, pdf_path)
        build_workbook(empty, xlsx_path)
        self.assertTrue(pdf_path.exists())
        workbook = load_workbook(xlsx_path)
        self.assertIn("Dashboard", workbook.sheetnames)
        pdf_path.unlink(missing_ok=True)
        xlsx_path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
